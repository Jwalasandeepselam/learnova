"""Source integrity checks for reusable file parsers; no production lesson seeds."""
import io
import pytest
from backend.app.ingestion.parsers import DocumentParser


def test_unknown_binary_and_legacy_files_are_not_reported_as_text():
    for name in ("program.exe", "old.doc", "old.ppt"):
        with pytest.raises(ValueError):
            DocumentParser.parse_file(b"arbitrary bytes", name)


def test_docx_tables_preserved_and_locations_identified_as_segments():
    from docx import Document
    doc = Document()
    doc.add_heading("Uploaded table", 1)
    row = doc.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = "Source value"
    row.cells[1].text = "4817"
    stream = io.BytesIO()
    doc.save(stream)
    parsed = DocumentParser.parse_file(stream.getvalue(), "table.docx")
    assert "4817" in parsed.raw_text
    assert parsed.metadata["page_numbers_are_estimates"] is True


def test_slide_table_is_included_in_its_source_slide():
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(5), Inches(1)).table
    table.cell(0, 0).text = "Source value"
    table.cell(0, 1).text = "9231"
    stream = io.BytesIO()
    deck.save(stream)
    parsed = DocumentParser.parse_file(stream.getvalue(), "table.pptx")
    assert "9231" in parsed.pages[0].text
    assert parsed.pages[0].page_number == 1
